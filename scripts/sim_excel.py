# -*- coding: utf-8 -*-
"""แปลงผลจำลอง (scripts/sim-orders.js → JSON) เป็น Excel สรุปแยกตามซัพและสาขา
   ใช้: python3 scripts/sim_excel.py sim.json out.xlsx [--label "ข้อมูล ณ ..."] [--verify verify.json]"""
import json,sys,argparse
from collections import defaultdict,OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
ap=argparse.ArgumentParser(); ap.add_argument('sim'); ap.add_argument('out')
ap.add_argument('--label',default=''); ap.add_argument('--verify',default='')
a=ap.parse_args()
D=json.load(open(a.sim,encoding='utf-8'))
VER=json.load(open(a.verify,encoding='utf-8')) if a.verify else None
BRN={'JJRD':'รัชดา','JJLP':'ลาดพร้าว'}
DOW_TH=['อาทิตย์','จันทร์','อังคาร','พุธ','พฤหัสบดี','ศุกร์','เสาร์']
DOW_EN={'sun':'อาทิตย์','mon':'จันทร์','tue':'อังคาร','wed':'พุธ','thu':'พฤหัสบดี','fri':'ศุกร์','sat':'เสาร์'}
ORD=['mon','tue','wed','thu','fri','sat','sun']
def th(ds):  # 2026-11-03 → "3 พ.ย. (อังคาร)"
    if not ds: return '-'
    import datetime
    d=datetime.date.fromisoformat(ds); m=['ม.ค.','ก.พ.','มี.ค.','เม.ย.','พ.ค.','มิ.ย.','ก.ค.','ส.ค.','ก.ย.','ต.ค.','พ.ย.','ธ.ค.'][d.month-1]
    return f"{d.day} {m} ({DOW_TH[(d.weekday()+1)%7]})"
def dayno(ds):  # เลขวันที่ในเดือน
    return int(ds[-2:])
def fq(v):
    if v is None: return ''
    return ('%.3f'%v).rstrip('0').rstrip('.')
def rule(s):
    if s['order_mode']=='fixed':
        sch=s.get('schedule') or {}
        parts=[f"สั่ง{DOW_EN[k]}→ส่ง{DOW_EN.get(sch[k],sch[k])}" for k in ORD if k in sch]
        t='วันสั่งตายตัว: '+(' · '.join(parts) if parts else '(ยังไม่ตั้งวัน!)')
        if s.get('order_ahead'): t+=f" · สั่งล่วงหน้า {s['order_ahead']} วัน"
    else:
        t=f"สั่งได้ทุกวัน · ส่งอีก {s.get('lead_days',1)} วัน"
    if s.get('prepay'): t+=' · 💸 จ่ายก่อนส่ง'
    return t
import datetime
wb=Workbook(); wb.remove(wb.active)
HEAD=Font(bold=True,color='FFFFFF'); FILL=PatternFill('solid',fgColor='7E1013'); SUB=PatternFill('solid',fgColor='FBF6EA')
WE=PatternFill('solid',fgColor='FFF3E0'); OK=PatternFill('solid',fgColor='EAF4EC'); BAD=PatternFill('solid',fgColor='FCEFEF')
thin=Side(style='thin',color='E5DCC8'); BORDER=Border(bottom=thin)
def sheet(title,cols,widths):
    ws=wb.create_sheet(title[:31]); ws.append(cols)
    for c in range(1,len(cols)+1):
        ws.cell(1,c).font=HEAD; ws.cell(1,c).fill=FILL
        ws.cell(1,c).alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width=widths[c-1]
    ws.freeze_panes='A2'; return ws
def wrap(ws,cols):
    for row in ws.iter_rows(min_row=2):
        for c in cols: row[c-1].alignment=Alignment(wrap_text=True,vertical='top')
sups=D['suppliers']; B=D['branches']
# ---------- สรุปกติกา + ผลจริงในช่วง ----------
def orders_of(code,sup):
    out=[]
    for day in B[code]['days']:
        g=next((x for x in day['plan'] if x['sup']==sup),None)
        if g and g['d']: out.append((day,g))
    return out
# ---------- ชีต "อ่านก่อน" ----------
ws=wb.create_sheet('อ่านก่อน')
n_days=len(B['JJRD']['days']); start=B['JJRD']['days'][0]['cd']; end=B['JJRD']['days'][-1]['cd']
for row in [
 ['ทดสอบจำลองปฏิทินสั่งของ · JJ เช็คสต๊อก'],[],
 ['ช่วงที่จำลอง',f"วันที่ 1–{n_days} · {th(start)} ถึง {th(end)}  (ใช้ปฏิทิน พ.ย. 2569 ที่วันที่ 1 ตรงกับวันอาทิตย์)"],
 ['ข้อมูลที่ใช้',a.label or D.get('generatedFrom','')],
 ['วิธีจำลอง','รันสูตรจริงของแอพ (หน้า 🛒 สั่งของ) ไล่วันขายทีละวัน ทั้ง 2 สาขา'],
 ['สมมติ 1','นับสต๊อกทุกวัน ได้ 0 ทุกตัว (ไม่มีของค้าง) → ยอดสั่ง = ยอดที่ต้องเผื่อจนของชุดหน้ามา'],
 ['สมมติ 2','ไม่มีใบสั่งค้างทาง (ไม่หัก "ของกำลังมา")'],
 ['สมมติ 3','สินค้าที่ยังไม่ตั้งอัตราใช้ = ไม่ขึ้นให้สั่ง (นับเป็น "ยังไม่ตั้งอัตรา") — ไม่ได้แปลว่าไม่ต้องสั่ง ดูคอลัมน์ "ยังไม่ตั้งอัตรา"'],
 ['สมมติ 4','เปิดร้านต่อเนื่อง 15 วัน ไม่คิดวันหยุดร้าน/วันหยุดซัพ/นักขัตฤกษ์ — ถ้ามีวันหยุด วันสั่ง–วันส่งจะเลื่อน'],
 ['ข้อควรรู้','ปริมาณในรายงาน = ของที่ต้องใช้ทั้งหมดในช่วงที่ของชุดนั้นต้องครอบคลุม ไม่ใช่ปริมาณที่จะสั่งจริง (ของจริงหักยอดนับได้ + ของที่สั่งค้างอยู่)'],
 ['ข้อควรรู้','กติกาซัพใช้ร่วมกันทั้ง 2 สาขา → วันสั่ง/วันส่งของทุกซัพเหมือนกันทั้งรัชดาและลาดพร้าว ต่างกันแค่รายการและปริมาณ (ชีต ซัพ × สาขา แสดงทั้ง 2 สาขาไว้เทียบกัน)'],
 ['ข้อควรรู้','ซัพที่ตั้ง "สั่งล่วงหน้า" ใบสั่งเดียวกันจะโผล่หลายวันติด — ในปฏิทินจะติดป้าย "(ใบเดียวกับวันก่อน)" และไม่นับซ้ำ'],
 ['ข้อควรรู้','คอลัมน์ที่ระบุ "(ระบบเดิม)" คือค่าที่แอพนับเดิมเคยใช้ หน้าสั่งของใหม่ยังไม่ใช้ค่านั้น'],[],
 ['ชีต','มีอะไร'],
 ['กติกาซัพ','กติกาที่ตั้งไว้ของแต่ละซัพ + สรุปว่าใน 15 วันสั่งกี่ครั้ง ส่งวันไหนบ้าง — ดูก่อนว่าตรงใจไหม'],
 ['ซัพ × สาขา','ทุกรอบสั่งของแต่ละซัพ แยกสาขา: สั่งวันไหน (สั่งจริงวันไหนถ้าตั้งล่วงหน้า) ส่งวันไหน ครอบคลุมถึงวันไหน กี่รายการ'],
 ['ปฏิทิน รัชดา / ลาดพร้าว','มองเป็นรายวัน: วันนี้ต้องสั่งซัพไหน · ของใครมาส่ง · ต้องนับสต๊อกอะไร'],
 ['เช็คสต๊อก','รายวัน: ต้องนับสินค้าของแผนก/ซัพไหนบ้าง (นับเฉพาะของที่จะสั่งวันนั้น)'],
 ['รายละเอียดสินค้า','ทุกบรรทัด: สาขา วัน ซัพ สินค้า หน่วย ยอดที่ต้องเผื่อ กี่วัน อัตราใช้'],
 ['ผลตรวจ','ตรวจสอบอัตโนมัติว่าเอนจินทำตามกติกา (วันส่งถูก · สั่งเฉพาะวันที่ตั้ง · ครอบคลุมถึงชุดหน้า)'],
]: ws.append(row)
ws['A1'].font=Font(bold=True,size=14,color='7E1013'); ws.column_dimensions['A'].width=26; ws.column_dimensions['B'].width=110
for r in (10,): 
    for c in (1,2): ws.cell(r,c).font=Font(bold=True)
# ---------- กติกาซัพ ----------
ws=sheet('กติกาซัพ',['ซัพ','กติกาที่ตั้งไว้','สินค้า รัชดา','สินค้า ลาดพร้าว','มีอัตราใช้ (รัชดา/ลาดพร้าว)','สั่งกี่ครั้งใน 15 วัน','วันที่สั่ง','วันที่ของมาส่ง','ของมาส่งวันหยุด (ยืนยันว่าซัพส่งจริง)','ผูกกลุ่มไลน์','เวลาตัดรอบ (ระบบเดิม)','รอบทุกกี่วัน (ระบบเดิม)','ขั้นต่ำ ลัง (ระบบเดิม)','⚠ ข้อสังเกต'],
  [30,52,10,10,16,12,34,34,16,10,12,12,12,60])
supnames=sorted(sups,key=lambda n:n.lower())
HAS_LG=any(x.get('line_group_id') for x in sups.values())   # สแนปช็อตเก่าไม่มีคอลัมน์กลุ่มไลน์ → อย่าเตือนมั่ว
def obs(s,od):   # ข้อสังเกตที่เจ้าของร้านควรเห็น
    o=[]
    if s['order_mode']=='interval':
        o.append("ตั้งเป็น 'interval' (รอบทุก N วัน ของระบบเดิม) — หน้าสั่งของใหม่ยังไม่รู้จักโหมดนี้ → ตอนนี้คิดเหมือน 'สั่งได้ทุกวัน ส่งอีก "+str(s.get('lead_days',1))+" วัน' (ขึ้นให้สั่งทุกวัน) ถ้าจริง ๆ สั่งเป็นรอบ ต้องตั้งเป็น 'วันสั่งตายตัว' เลือกวัน")
    if s['order_mode']=='any' and s.get('schedule'):
        o.append('ตั้งเป็น "สั่งได้ทุกวัน" แต่ยังมีตารางวันสั่ง→วันส่งค้างอยู่ (ระบบไม่ใช้ ถ้าจริง ๆ สั่งเป็นวัน ต้องสลับเป็น "วันสั่งตายตัว")')
    if s['order_mode']=='any' and (s.get('cycle_days') or 0)>1:
        o.append(f"ระบบเดิมตั้ง 'รอบทุก {int(s['cycle_days'])} วัน' แต่หน้าสั่งของใหม่ยังไม่ใช้ค่านี้ → ตอนนี้ขึ้นให้สั่งทุกวัน ถ้าจริง ๆ สั่งเป็นรอบ ต้องตั้งเป็น 'วันสั่งตายตัว'")
    if (s.get('min_cases') or 0)>0:
        o.append(f"ขั้นต่ำ {int(s['min_cases'])} ลัง ยังไม่ถูกใช้ปัดยอดสั่ง (รอเคาะ)")
    if s['order_mode']=='any' and s.get('lead_days',1)>1:
        o.append(f"ส่งอีก {s['lead_days']} วัน + สั่งได้ทุกวัน = สั่งทุกวันแต่ละใบเผื่อแค่ 1 วัน (ถ้าซัพส่งเป็นรอบจริง ควรตั้งวันสั่งตายตัว)")
    if s['order_mode']=='fixed' and not s.get('schedule'):
        o.append('ตั้งเป็นวันสั่งตายตัวแต่ยังไม่เลือกวัน → ไม่ขึ้นให้สั่งเลย')
    if s['order_mode']=='fixed' and s.get('order_ahead'):
        o.append(f"สั่งล่วงหน้า {s['order_ahead']} วัน: รายการจะโผล่ทั้งวันล่วงหน้าและวันสั่งจริง — วันสั่งจริงระบบหัก 'ของกำลังมา' จากใบสั่งที่ส่งไปแล้ว ถ้าวันล่วงหน้าไม่ได้กดส่งเข้าไลน์ (แค่ดู) ยอดวันจริงจะขึ้นเต็มอีกรอบ")
    if s['order_mode']=='fixed' and s.get('schedule'):
        from collections import defaultdict as _dd
        by=_dd(list)
        for k,v in s['schedule'].items(): by[v].append(k)
        for v,ks in by.items():
            if len(ks)>1:
                o.append('สั่ง'+'และ'.join(DOW_EN[k] for k in ORD if k in ks)+f" → ส่ง{DOW_EN.get(v,v)}วันเดียวกัน — ระบบจะขึ้นให้สั่งของชุดเดียวกันทั้ง 2 วัน (วันหลังจะหักเฉพาะใบที่กดส่งไลน์ไปแล้ว) ถ้าจริง ๆ สั่งวันเดียว ควรเอาวันหนึ่งออก")
    if s['order_mode']=='fixed' and od:
        mx=max(g['covDays'] for d,g in od); mn=min(g['covDays'] for d,g in od)
        if mx>=5: o.append(f"บางรอบต้องเผื่อถึง {mx} วัน (ของสดอาจไม่ไหว)")
        if mx-mn>=3: o.append(f"แต่ละรอบเผื่อไม่เท่ากัน ({mn}–{mx} วัน) ยอดสั่งจะแกว่งตามรอบ ปกติถ้าวันส่งห่างไม่เท่ากัน")
    if HAS_LG and not s.get('line_group_id'): o.append('ยังไม่ผูกกลุ่มไลน์ → ส่งใบสั่งอัตโนมัติไม่ได้')
    return o
for n in supnames:
    s=sups[n]
    cnt={c:sum(1 for day in B[c]['days'] for g in day['plan'] if g['sup']==n for _ in g['rows'][:1] for __ in [0]) for c in B}
    nprod={c:len(next((g['rows'] for g in B[c]['days'][0]['plan'] if g['sup']==n),[])) for c in B}
    rated={c:len([r for g in B[c]['days'][0]['plan'] if g['sup']==n for r in g['rows'] if (r['rate_wk'] or r['rate_fri'] or r['rate_we'])]) for c in B}
    od=orders_of('JJRD',n)
    wd=sorted({('เสาร์' if datetime.date.fromisoformat(g['d']).weekday()==5 else 'อาทิตย์') for d,g in od if datetime.date.fromisoformat(g['d']).weekday()>=5})
    ws.append([n,rule(s),nprod['JJRD'],nprod['JJLP'],f"{rated['JJRD']}/{rated['JJLP']}",len(od),
      ', '.join(str(dayno(d['cd'])) for d,g in od),', '.join(sorted({str(dayno(g['d'])) for d,g in od},key=int)),
      ', '.join(wd) if wd else '-',
      ('ผูกแล้ว' if s.get('line_group_id') else '⚠ ยังไม่ผูก') if HAS_LG else '(ไม่มีข้อมูล)',s.get('cutoff') or '',s.get('cycle_days') or '',s.get('min_cases') or '',
      ' · '.join(obs(s,od))])
    if s['order_mode']=='fixed': 
        for c in range(1,15): ws.cell(ws.max_row,c).fill=SUB
wrap(ws,[2,7,8,14]); ws.auto_filter.ref=ws.dimensions
# ---------- ซัพ × สาขา ----------
ws=sheet('ซัพ × สาขา',['ซัพ','สาขา','วันสั่ง','วันสั่งจริง (ถ้าล่วงหน้า)','วันส่ง','ส่งวันหยุด?','ของชุดหน้ามา','ต้องเผื่อกี่วัน','ครอบคลุมวัน','รายการที่ต้องสั่ง','ยังไม่ตั้งอัตรา','จ่ายก่อนส่ง','หมายเหตุ'],
  [28,10,18,20,18,10,18,10,26,12,10,10,46])
for n in supnames:
    s=sups[n]
    ws.append([n,'',rule(s)]+['']*10)
    for c in range(1,14): ws.cell(ws.max_row,c).fill=SUB; ws.cell(ws.max_row,c).font=Font(bold=True)
    for code in ('JJRD','JJLP'):
        od=orders_of(code,n)
        if not od:
            ws.append(['',BRN[code],'— ไม่มีรอบสั่งในช่วงนี้ —']+['']*10); continue
        seen=set()
        for day,g in od:
            note=[]
            dup=(g['od'] or g['d']) in seen; seen.add(g['od'] or g['d'])
            if dup: note.append('(ใบเดียวกับวันก่อน — สั่งล่วงหน้า/สองวันสั่งส่งวันเดียวกัน)')
            if g.get('unset'): note.append('⚠ ซัพยังไม่ตั้งรอบสั่ง (ระบบเดาว่าส่งพรุ่งนี้)')
            if g['nOrder']==0: note.append('ไม่มีอะไรต้องสั่ง (ยังไม่ตั้งอัตราใช้ทั้งหมด)')
            dd=__import__('datetime').date.fromisoformat(g['d']).weekday()
            ws.append([ '',BRN[code],th(day['cd']),th(g['od']) if g['ahead'] else '',th(g['d']),('ส่ง'+('เสาร์' if dd==5 else 'อาทิตย์')+' ⚠') if dd>=5 else '',
                th(g['nx']) if g['nx'] else '',
                g['covDays'],(g['covTxt'] or '').replace('ครอบคลุม ','').replace(' วัน (',' วัน: ').rstrip(')'),g['nOrder'],g['nNoRate'],
                '💸' if g['prepay'] else '',' · '.join(note)])
            if day['dow'] in (0,6) or dd>=5:
                for c in range(1,14): ws.cell(ws.max_row,c).fill=WE
            if dup:
                for c in range(1,14): ws.cell(ws.max_row,c).font=Font(color='8A8A8A')
wrap(ws,[3,4,5,7,9,13])
# ---------- ปฏิทินรายสาขา ----------
for code in ('JJRD','JJLP'):
    ws=sheet('ปฏิทิน '+BRN[code],['วันที่','วัน','ต้องสั่งกี่ซัพ','สั่งซัพไหน (รายการ · ส่งวันที่)','ของมาส่งวันนี้จาก','ต้องนับสต๊อก (แผนก: กี่รายการ)','💸 เตือนจ่ายก่อนส่ง','ซัพที่ยังไม่ผูกไลน์ (ส่งอัตโนมัติไม่ได้)'],
      [8,10,10,60,40,40,22,30])
    days=B[code]['days']
    deliver=defaultdict(list)
    for day in days:
        for g in day['plan']:
            if g['d'] and g['nOrder']>0: deliver[g['d']].append(f"{g['sup']} ({g['nOrder']}) จากใบสั่งวันที่ {dayno(day['cd'])}")
    seenOrd=set()   # (ซัพ, วันสั่งจริง) ที่โผล่ไปแล้ว — ใบเดียวกันไม่นับซ้ำ
    for day in days:
        ordering=[g for g in day['plan'] if g['d'] and g['nOrder']>0]
        fresh=[g for g in ordering if (g['sup'],g['od'] or g['d']) not in seenOrd]
        for g in ordering: seenOrd.add((g['sup'],g['od'] or g['d']))
        dept=defaultdict(int)
        for g in day['plan']:                       # นับสต๊อก = ของทุกซัพที่วันนี้เป็นวันสั่ง (แม้ยังไม่ตั้งอัตราก็ต้องนับ)
            if not g['d']: continue
            for r in g['rows']: dept[r['dept'] or 'ยังไม่จัดแผนก']+=1
        ws.append([dayno(day['cd']),DOW_TH[day['dow']],len(fresh),
          '\n'.join(f"{g['sup']} · {g['nOrder']} รายการ · ส่ง {dayno(g['d'])} {DOW_TH[int(__import__('datetime').date.fromisoformat(g['d']).strftime('%w'))]}"+(f" · สั่งล่วงหน้า {g['ahead']} วัน" if g['ahead'] else '')+('' if g in fresh else ' (ใบเดียวกับวันก่อน)') for g in ordering),
          '\n'.join(deliver.get(day['cd'],[])) or '—',
          '\n'.join(f"{k}: {v}" for k,v in sorted(dept.items(),key=lambda kv:-kv[1])),
          '\n'.join(g['sup'] for g in ordering if g['prepay']),
          ('\n'.join(g['sup'] for g in ordering if not g['lineGroup'])) if HAS_LG else '(ไม่มีข้อมูลในสแนปช็อตนี้)'])
        ws.row_dimensions[ws.max_row].height=max(18,15*max(1,len(ordering),len(deliver.get(day['cd'],[])),len(dept)))
        if day['dow'] in (0,6):
            for c in range(1,9): ws.cell(ws.max_row,c).fill=WE
    wrap(ws,[4,5,6,7,8])
# ---------- เช็คสต๊อก ----------
ws=sheet('เช็คสต๊อก',['วันที่','วัน','สาขา','ต้องนับกี่รายการ','แผนก: จำนวน','นับของซัพไหน (จำนวน)'],[8,10,10,12,40,70])
for code in ('JJRD','JJLP'):
    for day in B[code]['days']:
        ordering=[g for g in day['plan'] if g['d']]
        dept=defaultdict(int); n=0
        for g in ordering:
            for r in g['rows']: dept[r['dept'] or 'ยังไม่จัดแผนก']+=1; n+=1
        ws.append([dayno(day['cd']),DOW_TH[day['dow']],BRN[code],n,
          ', '.join(f"{k}: {v}" for k,v in sorted(dept.items(),key=lambda kv:-kv[1])),
          ', '.join(f"{g['sup']} ({len(g['rows'])})" for g in ordering)])
        if day['dow'] in (0,6):
            for c in range(1,7): ws.cell(ws.max_row,c).fill=WE
daily=[n for n,x in sups.items() if x['order_mode']!='fixed']
ws.append([]); ws.append(['สรุป','','','',
  ('ไม่มีวันไหนที่ไม่ต้องนับเลย — เพราะซัพโหมด "สั่งได้ทุกวัน/interval" '+str(len(daily))+' ราย ขึ้นให้สั่งทุกวัน: '+', '.join(daily)) if daily else 'มีวันที่ไม่ต้องนับอะไรเลย: ดูแถวที่จำนวน = 0',''])
ws.cell(ws.max_row,1).font=Font(bold=True)
wrap(ws,[5,6]); ws.auto_filter.ref='A1:F'+str(ws.max_row-2)
# ---------- รายละเอียดสินค้า ----------
ws=sheet('รายละเอียดสินค้า',['สาขา','วันที่','วัน','ซัพ','สินค้า','แผนก','หน่วยนับ','ต้องเผื่อ (หน่วย)','เผื่อกี่วัน','วันส่ง','ชุดหน้ามา','ใช้/วัน จ–พฤ','ศ','ส–อา','จ่ายก่อนส่ง'],
  [10,8,10,26,30,16,10,14,10,16,16,10,8,8,10])
for code in ('JJRD','JJLP'):
    for day in B[code]['days']:
        for g in day['plan']:
            if not g['d']: continue
            for r in g['rows']:
                if not r['need']: continue
                ws.append([BRN[code],dayno(day['cd']),DOW_TH[day['dow']],g['sup'],r['name'],r['dept'],r['unit'],r['need'],r['covDays'],th(g['d']),th(g['nx']) if g['nx'] else '',r['rate_wk'],r['rate_fri'],r['rate_we'],'💸' if g['prepay'] else ''])
ws.auto_filter.ref=ws.dimensions
for i in range(2,ws.max_row+1): ws.cell(i,8).number_format='0.###'
# ---------- ผลตรวจ ----------
ws=sheet('ผลตรวจ',['ซัพ','สาขา','ข้อที่ตรวจ','ผล','รายละเอียด'],[28,10,44,10,80])
import datetime
def check(sup,code,what,ok,detail=''):
    ws.append([sup,BRN.get(code,code),what,'✅ ผ่าน' if ok else '❌ ไม่ผ่าน',detail])
    for c in range(1,6): ws.cell(ws.max_row,c).fill=OK if ok else BAD
def dow_en(ds): return ORD[datetime.date.fromisoformat(ds).weekday()]
def addd(ds,n): return (datetime.date.fromisoformat(ds)+datetime.timedelta(days=n)).isoformat()
def days_between(a,b):
    out=[]; x=a
    while x<b: out.append(x); x=addd(x,1)
    return out
def rate_on(r,ds):
    wd=datetime.date.fromisoformat(ds).weekday()   # Mon=0..Sun=6
    v=r['rate_fri'] if wd==4 else r['rate_we'] if wd>=5 else r['rate_wk']
    return float(v or 0)
for n in supnames:
    s=sups[n]
    for code in ('JJRD','JJLP'):
        od=orders_of(code,n)
        if not od: continue
        # 1) วันส่งต้องหลังวันสั่ง
        bad=[(d['cd'],g['d']) for d,g in od if not g['d']>d['cd']]
        check(n,code,'วันส่งอยู่หลังวันที่กดสั่งเสมอ',not bad,'; '.join(f"{a}→{b}" for a,b in bad))
        if s['order_mode']=='fixed':
            sch=s.get('schedule') or {}
            # 2) วันสั่งจริงต้องเป็นวันที่ตั้งไว้ และวันส่งตรงตามคู่
            bad=[(d['cd'],g['od'],g['d']) for d,g in od if dow_en(g['od']) not in sch or dow_en(g['d'])!=sch[dow_en(g['od'])]]
            check(n,code,'สั่งเฉพาะวันที่ตั้ง และส่งตามวันคู่กัน (เช่น สั่งจันทร์→ส่งพุธ)',not bad,'; '.join(map(str,bad)))
            # 3) สั่งล่วงหน้า: วันกด = วันสั่งจริง − N
            ah=s.get('order_ahead') or 0
            bad=[(d['cd'],g['od'],g['ahead']) for d,g in od if g['ahead']>ah or (g['ahead']!=(datetime.date.fromisoformat(g['od'])-datetime.date.fromisoformat(d['cd'])).days)]
            check(n,code,f'สั่งล่วงหน้าไม่เกิน {ah} วัน และวันสั่งจริงอยู่ข้างหน้าตามที่ตั้ง',not bad,'; '.join(map(str,bad)))
            # 4) ทุกวันที่ตั้งไว้ในช่วง 15 วัน มีใบสั่ง (ไม่ตกหล่น)
            expect=[day['cd'] for day in B[code]['days'] if dow_en(day['cd']) in sch]
            got={g['od'] for d,g in od}
            miss=[x for x in expect if x not in got]
            check(n,code,'ทุกวันสั่งที่ตั้งไว้ในช่วงนี้มีใบสั่งครบ',not miss,'ตกหล่น: '+', '.join(miss) if miss else f'{len(expect)} รอบ')
            bad=[(d['cd'],g['d'],r['pre']) for d,g in od for r in g['rows'][:1] if r['need'] and r['pre'] is not None and round(r['pre'],3)!=round(sum(rate_on(r,x) for x in days_between(addd(d['cd'],1),g['d'])),3)]
            check(n,code,'หัก "ใช้ระหว่างรอของมาส่ง" ครบทุกวันที่รอ (รวมวันสั่งจริงกรณีสั่งล่วงหน้า)',not bad,'; '.join(map(str,bad[:4])))
            # 5) ครอบคลุมถึงชุดหน้า: วันส่ง + covDays = วันชุดหน้ามา
            bad=[(g['d'],g['covDays'],g['nx']) for d,g in od if g['nx'] and (datetime.date.fromisoformat(g['d'])+datetime.timedelta(days=g['covDays'])).isoformat()!=g['nx']]
            check(n,code,'ยอดสั่งเผื่อตั้งแต่วันของมา จนถึงก่อนของชุดถัดไปมา (ไม่มีวันโหว่/ไม่ซ้อน)',not bad,'; '.join(map(str,bad)))
        else:
            lead=s.get('lead_days',1)
            bad=[(d['cd'],g['d']) for d,g in od if (datetime.date.fromisoformat(g['d'])-datetime.date.fromisoformat(d['cd'])).days!=lead]
            tag='(interval → ระบบคิดเป็นสั่งได้ทุกวัน) ' if s['order_mode']=='interval' else ''
            check(n,code,f'{tag}สั่งได้ทุกวัน: ส่งหลังสั่ง {lead} วันทุกครั้ง',not bad,'; '.join(map(str,bad)))
            check(n,code,f'{tag}สั่งทุกวัน (15 รอบใน 15 วัน) และแต่ละรอบเผื่อ 1 วัน',len(od)==len(B[code]['days']) and all(g['covDays']==1 for d,g in od),f'{len(od)} รอบ')
            # ใช้ระหว่างรอของมาส่ง = ผลรวมอัตราใช้ของวันถัดจากวันกด ถึงก่อนวันส่ง (เอนจินหักออกจากยอดนับ)
            bad=[(d['cd'],g['d'],r['pre']) for d,g in od for r in g['rows'][:1] if r['need'] and r['pre'] is not None and round(r['pre'],3)!=round(sum(rate_on(r,x) for x in days_between(addd(d['cd'],1),g['d'])),3)]
            check(n,code,'หัก "ใช้ระหว่างรอของมาส่ง" ครบทุกวันที่รอ',not bad,'; '.join(map(str,bad[:4])))
        # 6) จ่ายก่อนส่ง
        if s.get('prepay'):
            check(n,code,'ซัพจ่ายก่อนส่ง → ทุกใบสั่งติดธง 💸',all(g['prepay'] for d,g in od))
for code in ('JJRD','JJLP'):
    rows=[(day,g) for day in B[code]['days'] for g in day['plan'] if g['sup']=='(ไม่ระบุซัพ)']
    if rows:
        bad=[d['cd'] for d,g in rows if not (g.get('unset') and g['d']==addd(d['cd'],1))]
        check('(ไม่ระบุซัพ)',code,'สินค้าที่ไม่มีซัพ/ซัพไม่อยู่ในตาราง → ระบบติดธง "ยังไม่ตั้งรอบสั่ง" และเดาว่าส่งพรุ่งนี้ ทุกวัน',not bad,
              f"{len(rows[0][1]['rows'])} รายการ: "+', '.join(r['name'] for r in rows[0][1]['rows']))
if VER:
    ws.append([]); ws.append(['ผลตรวจอิสระ (ผู้ตรวจคนละคน คิดมือจากกติกาแล้วเทียบ)']); ws.cell(ws.max_row,1).font=Font(bold=True)
    for v in VER:
        ws.append([v.get('sup',''),v.get('branch',''),v.get('check',''),'✅ ตรง' if v.get('ok') else '❌ ต่าง',v.get('detail','')])
        for c in range(1,6): ws.cell(ws.max_row,c).fill=OK if v.get('ok') else BAD
wrap(ws,[3,5])
wb.save(a.out); print('saved',a.out)
