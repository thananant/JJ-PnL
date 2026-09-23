# -*- coding: utf-8 -*-
"""คิด "อัตราใช้ต่อวัน" ของระบบนับสต๊อก จากไฟล์ export ของ P&L (แท็บรายสินค้า) หลายเดือนรวมกัน

ใช้:
  python3 scripts/rates_from_pnl.py --branch JJRD \
      --products ไฟล์ที่ได้จาก sql/jjmk_export_products.sql (CSV) \
      --out-sql sql/jjmk_stockcheck_rates_jjrd.sql \
      --out-xlsx /tmp/review.xlsx \
      ไฟล์เดือน1.xlsx ไฟล์เดือน2.xlsx ...

หลักการ: ปริมาณที่ซื้อของแต่ละกลุ่มวัน (จ–พฤ / ศ / ส–อา) รวมทุกเดือน ÷ จำนวนวันขายของกลุ่มนั้นรวมทุกเดือน
  ตัวเลขที่ได้อยู่ใน "หน่วยที่ลงบิล" → ปล่อยให้ SQL แปลงเป็นหน่วยนับเองตามตัวคูณที่ตั้งไว้ในฐานข้อมูล
  สินค้าที่หน่วยบิล ≠ หน่วยนับ และยังไม่มีตัวคูณ = ข้ามไว้ ไม่เดาให้
"""
import argparse,csv,json,sys
from collections import defaultdict
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter

BR={'JJRD':('b19f0a17b4472','รัชดา'),'JJLP':('b19f0a17b448212','ลาดพร้าว')}
ALIAS={'กก.':'kg','กก':'kg','กิโลกรัม':'kg','กิโล':'kg','โล':'kg','kg':'kg','แพค':'แพ็ค'}
nz=lambda u:ALIAS.get((u or '').strip(),(u or '').strip().lower())
q=lambda s:"'"+str(s).replace("'","''")+"'"
NORM=("case when lower(btrim({0})) in ('กก.','กก','กิโลกรัม','กิโล','โล','kg') then 'kg'\n"
      "       when lower(btrim({0})) in ('แพ็ค','แพค','pack') then 'แพ็ค'\n"
      "       else lower(btrim({0})) end")

def read_months(paths):
    """คืน (cell, unlinked, months) — cell[(ชื่อนับ, หน่วย normalize)] = ยอดสะสม"""
    cell=defaultdict(lambda:{'wk':0.0,'fri':0.0,'we':0.0,'val':0.0,'qty':0.0,'lines':0,'unit':'','bill':set(),'sup':set()})
    unl=defaultdict(lambda:{'val':0.0,'units':set(),'sup':set()})
    months={}
    for p in paths:
        wb=openpyxl.load_workbook(p,data_only=True)
        s=list(wb['สรุป'].iter_rows(min_row=2,max_row=2,values_only=True))[0]
        months[s[1]]={'sales':s[2],'days':s[3],'wk':s[5],'fri':s[7],'we':s[9],'branch':s[0]}
        ws=wb['รายสินค้า']
        hdr=[c for c in next(ws.iter_rows(min_row=1,max_row=1,values_only=True))]
        ix={h:i for i,h in enumerate(hdr)}
        wkd,frd,wed=s[5],s[7],s[9]
        for r in ws.iter_rows(min_row=2,values_only=True):
            if not r[0]:continue
            nm=(r[ix['ชื่อนับ']] or '').strip(); u=str(r[ix['หน่วย']] or '').strip()
            if not nm:
                x=unl[str(r[0]).strip()]; x['val']+=float(r[ix['มูลค่า (บาท)']] or 0); x['units'].add(u)
                x['sup'].add(str(r[ix['ซัพพลายเออร์']] or '').strip()); continue
            d=cell[(nm,nz(u))]; d['unit']=u
            d['bill'].add(str(r[0]).strip()); d['sup'].add(str(r[ix['ซัพพลายเออร์']] or '').strip())
            d['val']+=float(r[ix['มูลค่า (บาท)']] or 0); d['qty']+=float(r[ix['ปริมาณรวม']] or 0)
            d['lines']+=int(r[ix['จำนวนบรรทัดบิล']] or 0)
            d['wk'] +=float(r[ix['จ–พฤ ใช้/วัน']] or 0)*wkd
            d['fri']+=float(r[ix['ศ ใช้/วัน']]   or 0)*frd
            d['we'] +=float(r[ix['ส–อา ใช้/วัน']] or 0)*wed
    return cell,unl,months

def build(cell,months,prod):
    TW=sum(m['wk'] for m in months.values()); TF=sum(m['fri'] for m in months.values()); TE=sum(m['we'] for m in months.values())
    urows=[]
    for (nm,un),d in sorted(cell.items()):
        urows.append({'name':nm,'unit':d['unit'],'nunit':un,
            'wk':round(d['wk']/TW,3),'fri':round(d['fri']/TF,3),'we':round(d['we']/TE,3),
            'val':round(d['val']),'qty':round(d['qty'],2),'lines':d['lines'],
            'price':round(d['val']/d['qty'],2) if d['qty'] else None,
            'sup':' / '.join(sorted(x for x in d['sup'] if x))})
    byname=defaultdict(list)
    for r in urows: byname[r['name']].append(r)
    f2=lambda v:(None if v in (None,'','null') else float(v))
    ap,hold=[],[]
    for nm,lst in byname.items():
        p=prod.get(nm)
        if not p: continue
        cu=nz(p['หน่วยนับ'])
        base={'name':nm,'countUnit':p['หน่วยนับ'],'dept':p['แผนก'],'sup':p['ชื่อซัพ'],
              'val':sum(r['val'] for r in lst),'lines':sum(r['lines'] for r in lst),
              'curwk':f2(p['ใช้ต่อวัน จ-พฤ']),'curfri':f2(p['ใช้ต่อวัน ศุกร์']),'curwe':f2(p['ใช้ต่อวัน ส-อา']),
              'units':'; '.join('{} {:,.2f} ({:,.0f} บาท)'.format(r['unit'],r['qty'],r['val']) for r in sorted(lst,key=lambda x:-x['val']))}
        if all(r['nunit']==cu for r in lst):
            ap.append({**base,'wk':round(sum(r['wk'] for r in lst),3),'fri':round(sum(r['fri'] for r in lst),3),
                       'we':round(sum(r['we'] for r in lst),3)})
        else:
            main=max(lst,key=lambda x:x['val'])
            same=next((r for r in lst if r['nunit']==cu),None)
            sug=round(main['price']/same['price'],2) if (same and same is not main and same['price'] and main['price']) else None
            hold.append({**base,'billUnits':', '.join(sorted({r['unit'] for r in lst})),'suggest':sug,
                         'rateIn':main['unit'],'wk':main['wk'],'fri':main['fri'],'we':main['we']})
    nob=[{'name':k,'countUnit':v['หน่วยนับ'],'dept':v['แผนก'],'sup':v['ชื่อซัพ'],
          'curwk':f2(v['ใช้ต่อวัน จ-พฤ']),'curfri':f2(v['ใช้ต่อวัน ศุกร์']),'curwe':f2(v['ใช้ต่อวัน ส-อา'])}
         for k,v in prod.items() if k not in byname]
    ap.sort(key=lambda x:-x['val']); hold.sort(key=lambda x:-x['val']); nob.sort(key=lambda x:x['name'])
    return urows,ap,hold,nob,(TW,TF,TE)

def write_sql(path,code,urows,months,days,batch):
    bid,brname=BR[code]
    TW,TF,TE=days
    vals=",\n  ".join('({},{},{},{},{})'.format(q(r['name']),q(r['unit']),r['wk'],r['fri'],r['we']) for r in urows)
    mtxt=" · ".join('{}: ขาย {:,.0f} บาท {} วัน (จ–พฤ {} · ศ {} · ส–อา {})'.format(
        m,v['sales'],v['days'],v['wk'],v['fri'],v['we']) for m,v in sorted(months.items()))
    sql=open(__file__.replace('rates_from_pnl.py','rates_template.sql'),encoding='utf-8').read()
    sql=(sql.replace('@@BID@@',bid).replace('@@BR@@',brname).replace('@@CODE@@',code)
            .replace('@@BATCH@@',batch).replace('@@MONTHS@@',mtxt)
            .replace('@@DW@@',str(TW)).replace('@@DF@@',str(TF)).replace('@@DE@@',str(TE))
            .replace('@@VALUES@@',vals)
            .replace('@@N_BILLUNIT@@',NORM.format('bill_unit')).replace('@@N_PUNIT@@',NORM.format('p.unit'))
            .replace('@@N_NBILL@@',NORM.format('n.bill_unit')).replace('@@N_NSTOCK@@',NORM.format('n.stock_unit'))
            .replace('@@N_PPACK@@',NORM.format('p.pack_unit'))
            .replace('@@N_CFROM@@',NORM.format('c.from_unit')).replace('@@N_CTO@@',NORM.format('c.to_unit')))
    open(path,'w',encoding='utf-8').write(sql)

def write_xlsx(path,ap,hold,nob,unl,months,days):
    wb=Workbook(); wb.remove(wb.active)
    HEAD=Font(bold=True,color='FFFFFF'); FILL=PatternFill('solid',fgColor='7E1013')
    def sheet(title,cols,rows,widths,fmts=None):
        ws=wb.create_sheet(title[:31]); ws.append(cols)
        for c in range(1,len(cols)+1):
            ws.cell(1,c).font=HEAD; ws.cell(1,c).fill=FILL
            ws.cell(1,c).alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            ws.column_dimensions[get_column_letter(c)].width=widths[c-1]
        for r in rows: ws.append(r)
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for c,f in (fmts or {}).items():
            for i in range(2,ws.max_row+1): ws.cell(i,c).number_format=f
    rat=lambda n,o:(round(n/o,2) if o else None)
    conf=lambda n:('สูง' if n>=8 else 'กลาง' if n>=4 else 'ต่ำ — ซื้อไม่กี่ครั้ง ดูเอง')
    sheet('ใส่ให้แล้ว',
      ['สินค้า (ชื่อนับ)','แผนก','ซัพ','หน่วยนับ','ใหม่ จ–พฤ/วัน','ใหม่ ศุกร์/วัน','ใหม่ ส–อา/วัน',
       'เดิม จ–พฤ','เดิม ศุกร์','เดิม ส–อา','เปลี่ยนกี่เท่า (จ–พฤ)','ความมั่นใจ','มูลค่าบิล','บรรทัดบิล','หน่วยที่ลงบิล'],
      [[r['name'],r['dept'],r['sup'],r['countUnit'],r['wk'],r['fri'],r['we'],r['curwk'],r['curfri'],r['curwe'],
        rat(r['wk'],r['curwk']),conf(r['lines']),r['val'],r['lines'],r['units']] for r in ap],
      [28,16,22,10,13,13,13,11,11,11,15,20,15,10,42],
      {5:'0.000',6:'0.000',7:'0.000',8:'0.00',9:'0.00',10:'0.00',11:'0.00',13:'#,##0'})
    sheet('รอตั้งตัวคูณหน่วย',
      ['สินค้า (ชื่อนับ)','แผนก','ซัพ','หน่วยนับ (ที่นับจริง)','หน่วยที่ลงบิล','ต่อวัน (หน่วยบิล) จ–พฤ','ศุกร์','ส–อา',
       'หน่วยของตัวเลข','เดาตัวคูณจากราคา (1 หน่วยใหญ่ = กี่หน่วยนับ)','มูลค่าบิล','รายละเอียดหน่วย'],
      [[r['name'],r['dept'],r['sup'],r['countUnit'],r['billUnits'],r['wk'],r['fri'],r['we'],r['rateIn'],
        r['suggest'],r['val'],r['units']] for r in hold],
      [28,16,22,16,16,22,11,11,13,30,15,46],
      {6:'0.000',7:'0.000',8:'0.000',10:'0.00',11:'#,##0'})
    sheet('ช่วงนี้ไม่มีบิล',
      ['สินค้า (ชื่อนับ)','แผนก','ซัพ','หน่วยนับ','อัตราเดิม จ–พฤ','เดิม ศุกร์','เดิม ส–อา'],
      [[r['name'],r['dept'],r['sup'],r['countUnit'],r['curwk'],r['curfri'],r['curwe']] for r in nob],
      [28,16,22,10,14,12,12],{5:'0.00',6:'0.00',7:'0.00'})
    sheet('บิลที่ยังไม่ผูกชื่อนับ',['ชื่อในบิล (P&L)','หน่วย','มูลค่า','ซัพ'],
      [[k,', '.join(sorted(v['units'])),round(v['val']),' / '.join(sorted(x for x in v['sup'] if x))]
       for k,v in sorted(unl.items(),key=lambda kv:-kv[1]['val'])],[34,14,16,30],{3:'#,##0'})
    ws=wb.create_sheet('วิธีคิด',0)
    for row in [['อัตราใช้ต่อวัน — คำนวณจากบิลจริง'],[],
      ['เดือน','ยอดขายรวม VAT','วันขาย','วัน จ–พฤ','วัน ศ','วัน ส–อา'],
      *[[m,v['sales'],v['days'],v['wk'],v['fri'],v['we']] for m,v in sorted(months.items())],
      ['รวม','','',days[0],days[1],days[2]],[],
      ['สูตร','ปริมาณที่ซื้อทุกเดือนรวมกัน → เกลี่ยตามยอดขายของกลุ่มวัน → หารจำนวนวันขายในกลุ่มนั้น'],[],
      ['สรุป','จำนวน'],['ใส่ให้แล้ว',len(ap)],['รอตั้งตัวคูณหน่วยก่อน',len(hold)],
      ['ช่วงนี้ไม่มีบิล (คงค่าเดิม)',len(nob)],['บิลที่ยังไม่ผูกชื่อนับ',len(unl)],[],
      ['ข้อควรรู้','เป็นตัวเลข "ของที่ซื้อเข้า" ไม่ใช่ของที่ใช้จริงเป๊ะ ๆ — เดือนที่ตุนของเยอะจะสูงกว่าความจริง'],
      ['','ค่าเดิมสำรองไว้ที่ตาราง jjsc_rate_backup ย้อนกลับได้']]:
        ws.append(row)
    ws['A1'].font=Font(bold=True,size=14,color='7E1013')
    for c,w in zip('ABCDEF',[46,26,12,12,10,12]): ws.column_dimensions[c].width=w
    wb.save(path)

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('months',nargs='+')
    ap.add_argument('--branch',default='JJRD',choices=list(BR))
    ap.add_argument('--products',required=True,help='CSV จาก sql/jjmk_export_products.sql')
    ap.add_argument('--out-sql',required=True); ap.add_argument('--out-xlsx')
    a=ap.parse_args()
    brname=BR[a.branch][1]
    cell,unl,months=read_months(a.months)
    prod={}
    for r in csv.DictReader(open(a.products,encoding='utf-8')):
        if r['สาขา']==brname: prod.setdefault(r['ชื่อสินค้า (ชื่อนับ)'].strip(),r)
    urows,apply_,hold,nob,days=build(cell,months,prod)
    batch='rates_{}_{}'.format(a.branch.lower(),'+'.join(sorted(months)))
    write_sql(a.out_sql,a.branch,urows,months,days,batch)
    if a.out_xlsx: write_xlsx(a.out_xlsx,apply_,hold,nob,unl,months,days)
    print('เดือนที่ใช้:',', '.join(sorted(months)),'· วันขายรวม จ–พฤ {} · ศ {} · ส–อา {}'.format(*days))
    print('แถวหน่วย {} · ใส่ได้ {} · รอตัวคูณ {} · ไม่มีบิล {} · บิลยังไม่ผูกชื่อ {}'.format(
        len(urows),len(apply_),len(hold),len(nob),len(unl)))
    print('SQL →',a.out_sql, ('· Excel → '+a.out_xlsx) if a.out_xlsx else '')

if __name__=='__main__': main()
